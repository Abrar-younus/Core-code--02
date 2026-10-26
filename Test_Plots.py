import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook

file_path = r"C:\Users\M Muneeb ur Rehman\Desktop\FYP\NewNew\arslan\Tracing\trial_5_(star)_unfil.xlsx"

# Load workbook and extract data before "Demographics:"
wb = load_workbook(file_path)
ws = wb.active

data_rows = []
for row in ws.iter_rows(values_only=True):
    if row and any(str(cell).strip().lower() == "demographics:" for cell in row):
        break
    data_rows.append(row)

# Convert to DataFrame and clean it
df = pd.DataFrame([row for row in data_rows if any(row)])
df.columns = df.iloc[0]
df = df[1:]

# Extract coordinates
Original_X = df['Original_X'].dropna().astype(float).tolist()
Original_Y = df['Original_Y'].dropna().astype(float).tolist()
Traced_X = df['Traced_X'].dropna().astype(float).tolist()
Traced_Y = df['Traced_Y'].dropna().astype(float).tolist()

# Plotting
plt.figure(figsize=(10, 6))
plt.plot(Original_X, Original_Y, label='Original Path', color='red')
plt.plot(Traced_X, Traced_Y, label='Traced Path', color='blue')
plt.scatter(Original_X, Original_Y, color='red', s=5)
plt.scatter(Traced_X, Traced_Y, color='blue', s=5)

plt.title('Original and Traced Paths')
plt.xlabel('X')
plt.ylabel('Y')
plt.legend()
plt.axis('equal')
plt.grid(True)
plt.tight_layout()
plt.show()